import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.classifier import classify
from app.models import FontChoice, PaperType, RenderOptions
from app.parser import parse_docx
from app.renderer import render_docx


def _render(paper_bytes, **options):
    parsed = parse_docx(paper_bytes)
    classify(parsed.blocks, use_ai=False)
    return Document(io.BytesIO(render_docx(parsed, RenderOptions(**options))))


def _para(doc, text):
    return next(p for p in doc.paragraphs if p.text == text)


def test_page_setup_and_fonts(paper_bytes):
    doc = _render(paper_bytes)
    section = doc.sections[0]
    assert all(m.inches == 1 for m in (section.top_margin, section.bottom_margin,
                                        section.left_margin, section.right_margin))
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE
    # no direct font overrides survive on runs
    body = doc.element.body
    assert not list(body.iter(qn("w:rFonts"))) and not list(body.iter(qn("w:sz")))


def test_structure(paper_bytes):
    doc = _render(paper_bytes)
    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts[0] == "The Effects of Sleep on Memory Consolidation"
    assert texts.index("Abstract") < texts.index("Introduction") < texts.index("References")
    assert texts.count("The Effects of Sleep on Memory Consolidation") == 2  # title page + first text page

    intro = _para(doc, "Introduction")
    assert intro.style.name == "Heading 1"
    assert doc.styles["Heading 1"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert _para(doc, "Prior Work on Sleep").style.name == "Heading 2"
    assert _para(doc, "Theoretical Framing").style.name == "Heading 3"

    keywords = next(p for p in doc.paragraphs if p.text.startswith("Keywords:"))
    assert keywords.runs[0].italic and keywords.text == "Keywords: sleep, memory, consolidation"

    # caption moved above the table, split into bold label and italic title
    assert texts.index("Table 1") + 1 == texts.index("Participant Demographics by Condition")
    assert _para(doc, "Note. Ages are means in years.").runs[0].italic


def test_references_sorted_unnumbered_hanging(paper_bytes):
    doc = _render(paper_bytes)
    refs = [p for p in doc.paragraphs if p.style.name == "APA Reference"]
    assert [p.text[:5] for p in refs] == ["Alpha", "Zeta,"]
    style = doc.styles["APA Reference"].paragraph_format
    assert style.left_indent.inches == 0.5 and style.first_line_indent.inches == -0.5
    assert _para(doc, "References").paragraph_format.page_break_before


def test_content_preserved(paper_bytes):
    doc = _render(paper_bytes)
    reltypes = [r.reltype for r in doc.part.rels.values()]
    assert RT.IMAGE in reltypes and RT.HYPERLINK in reltypes
    body = doc.element.body
    assert list(body.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))
    assert len(doc.tables) == 1
    borders = doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders.find(qn("w:insideV")).get(qn("w:val")) == "nil"
    lists = [p for p in doc.paragraphs if p.style.name == "APA List"]
    assert len(lists) == 3 and all(p._p.pPr.numPr is not None for p in lists)
    drosophila = next(r for p in doc.paragraphs for r in p.runs if r.text == "Drosophila")
    assert drosophila.italic


def test_header_variants(paper_bytes):
    student = _render(paper_bytes)
    header = student.sections[0].header.paragraphs[0]
    assert header._p.find(qn("w:fldSimple")).get(qn("w:instr")) == "PAGE"

    pro = _render(paper_bytes, paper_type=PaperType.PROFESSIONAL, font=FontChoice.CALIBRI_11)
    assert pro.sections[0].header.paragraphs[0].text.startswith("THE EFFECTS OF SLEEP ON MEMORY CONSOLIDATION\t")
    assert pro.styles["Normal"].font.name == "Calibri"
    texts = [p.text for p in pro.paragraphs]
    assert "PSY 301: Research Methods" not in texts  # course info is student-paper only


def test_without_title_page(paper_bytes):
    doc = _render(paper_bytes, include_title_page=False)
    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts[0] == "Abstract"
    assert "Jane Q. Student" not in texts
