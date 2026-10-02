"""Build deliberately messy research papers for tests and manual tryouts."""

from __future__ import annotations

import io
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def tiny_png(width: int = 4, height: int = 3) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\x80\x40\x20" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _para(doc, text, *, bold=False, italic=False, size=None, center=False, font="Arial", color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold or None, italic or None
    run.font.name = font
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    return p


def _hyperlink(paragraph, url, text):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def messy_paper() -> bytes:
    doc = Document()
    _para(doc, "THE EFFECTS OF SLEEP ON MEMORY CONSOLIDATION", bold=True, size=16, center=True)
    _para(doc, "Jane Q. Student", center=True)
    _para(doc, "Department of Psychology, State University", center=True)
    _para(doc, "PSY 301: Research Methods", center=True)
    _para(doc, "Dr. Alan Turing", center=True)
    _para(doc, "March 3, 2026", center=True)
    doc.add_paragraph("")
    _para(doc, "Abstract", bold=True, size=14)
    _para(doc, "We tested whether sleep improves recall. Participants (N = 40) learned word pairs and were "
               "tested after sleep or wake. Recall was higher after sleep.")
    _para(doc, "Keywords: sleep, memory, consolidation")

    _para(doc, "1. Introduction", bold=True, size=14, color=(0x1F, 0x4E, 0x79))
    p = _para(doc, "\tSleep has long been thought to support memory. ")
    p.add_run("Drosophila").italic = True
    p.add_run(" studies agree (see ")
    _hyperlink(p, "https://example.org/sleep", "this review")
    p.add_run(").")
    _para(doc, "1.1 prior work on sleep", bold=True)
    _para(doc, "Earlier studies used small samples and short retention intervals.")
    _para(doc, "theoretical framing", bold=True, italic=True)
    _para(doc, "Active systems consolidation predicts a benefit of slow-wave sleep.")

    table = doc.add_table(rows=3, cols=3)
    table.style = doc.styles["Table Grid"]
    for r, row in enumerate([["Group", "N", "Age"], ["Sleep", "20", "21.3"], ["Wake", "20", "20.9"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    _para(doc, "Table 1. participant demographics by condition")
    _para(doc, "Note. Ages are means in years.")

    doc.add_picture(io.BytesIO(tiny_png()), width=Inches(2))
    _para(doc, "Figure 1: Recall accuracy in the sleep and wake groups")

    eq = doc.add_paragraph()
    eq._p.append(parse_xml(f'<m:oMathPara xmlns:m="{M_NS}"><m:oMath><m:r><m:t>d=(M1-M2)/SD</m:t></m:r>'
                           f"</m:oMath></m:oMathPara>"))

    _para(doc, "2. Method", bold=True, size=14)
    for item in ("Recruit participants", "Train word pairs", "Test recall"):
        doc.add_paragraph(item, style="List Number")
    _para(doc, "Participants were undergraduates who received course credit.")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _para(doc, "References", bold=True)
    _para(doc, "[1] Zeta, A. (2020). Sleep and memory. Journal of Sleep Research, 29(2), 1–10.")
    _para(doc, "[2] Alpha, B. (2019). Consolidation revisited. Academic Press.")
    _para(doc, "Appendix A: survey items")
    _para(doc, "Item 1. How many hours did you sleep last night?")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
