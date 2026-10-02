"""Build an APA 7 document from labeled blocks.

Source paragraphs and tables are deep-copied into a fresh document, so images, equations,
hyperlinks, footnotes and inline emphasis are preserved; only layout and fonts are reset.
"""

from __future__ import annotations

import copy
import io
import re

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from . import apa7
from .models import ABSTRACT_LABELS, HEADING_LEVELS, TITLE_PAGE_LABELS, Block, Label, PaperType, RenderOptions
from .parser import ParsedDocument, paragraph_text

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Run properties that set fonts, sizes and colours; APA wants one uniform font.
_RUN_LOOK = ("w:rFonts", "w:sz", "w:szCs", "w:color", "w:highlight", "w:shd", "w:kern", "w:w", "w:spacing",
             "w:effect", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:bdr")
_RUN_EMPHASIS = ("w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:u")


class _RelationshipCopier:
    """Re-point r:id / r:embed attributes of copied XML at relationships in the new part."""

    def __init__(self, src_part, dst_part):
        self.src_part = src_part
        self.dst_part = dst_part
        self.cache: dict[str, str] = {}

    def copy(self, el):
        new = copy.deepcopy(el)
        for node in new.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith(f"{{{R_NS}}}"):
                    new_id = self._map(value)
                    if new_id is None:
                        del node.attrib[attr]
                    else:
                        node.set(attr, new_id)
        return new

    def _map(self, rid: str) -> str | None:
        if rid in self.cache:
            return self.cache[rid]
        rel = self.src_part.rels.get(rid)
        if rel is None:
            return None
        if rel.is_external:
            new_id = self.dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_id = self.dst_part.relate_to(rel.target_part, rel.reltype)
        self.cache[rid] = new_id
        return new_id


class _Renderer:
    def __init__(self, parsed: ParsedDocument, options: RenderOptions):
        self.parsed = parsed
        self.options = options
        self.out: DocxDocument = Document()
        body = self.out.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        self.sect_pr = body.find(qn("w:sectPr"))
        apa7.setup_document(self.out, options.font)
        self.copier = _RelationshipCopier(parsed.doc.part, self.out.part)
        self._share_document_parts()
        self.style_ids = {s.style_id for s in self.out.styles}
        self.pending_page_break = False

    # --- document plumbing ------------------------------------------------------------

    def _share_document_parts(self) -> None:
        """Bring along the parts that copied paragraphs point into (numbering, footnotes)."""
        src_rels = {rel.reltype: rel for rel in self.parsed.doc.part.rels.values() if not rel.is_external}
        if RT.NUMBERING in src_rels:
            for rel in self.out.part.rels.values():
                if rel.reltype == RT.NUMBERING:
                    rel._target = src_rels[RT.NUMBERING].target_part
        for reltype in (RT.FOOTNOTES, RT.ENDNOTES):
            if reltype in src_rels:
                part = src_rels[reltype].target_part
                self.out.part.relate_to(part, reltype)
                for el in part.element.iter(qn("w:p")):
                    self._clean_runs(el, emphasis=False)

    def _append(self, el) -> None:
        if el.tag == qn("w:p") and self.pending_page_break:
            Paragraph(el, self.out._body).paragraph_format.page_break_before = True
            self.pending_page_break = False
        elif el.tag == qn("w:tbl") and self.pending_page_break:
            spacer = OxmlElement("w:p")
            Paragraph(spacer, self.out._body).paragraph_format.page_break_before = True
            self.sect_pr.addprevious(spacer)
            self.pending_page_break = False
        self.sect_pr.addprevious(el)

    def _paragraph(self, el) -> Paragraph:
        return Paragraph(el, self.out._body)

    def _new_paragraph(self, text: str, style: str, bold: bool | None = None, italic: bool | None = None):
        p = OxmlElement("w:p")
        para = self._paragraph(p)
        para.style = self.out.styles[style]
        if text:
            run = para.add_run(text)
            if bold is not None:
                run.bold = bold
            if italic is not None:
                run.italic = italic
        return p

    def _copy_paragraph(self, block: Block, style: str, *, keep_numbering: bool = False,
                        emphasis: bool = True, keep_italic: bool = True):
        src = self.parsed.elements[block.index]
        el = self.copier.copy(src)
        if keep_numbering:
            _inline_style_numbering(src, el, self.parsed.doc)
        self._clean_paragraph(el, style, keep_numbering=keep_numbering)
        self._clean_runs(el, emphasis=emphasis, keep_italic=keep_italic)
        _strip_leading_whitespace(el)
        return el

    def _clean_paragraph(self, el, style: str | None, *, keep_numbering=False, keep_alignment=False) -> None:
        ppr = el.find(qn("w:pPr"))
        if ppr is not None:
            for child in list(ppr):
                keep = (child.tag == qn("w:numPr") and keep_numbering) or (
                    child.tag == qn("w:jc") and keep_alignment)
                if not keep:
                    ppr.remove(child)
        if style is not None:
            self._paragraph(el).style = self.out.styles[style]
        # Manual page breaks and Word's cached layout hints have no place in the new layout.
        for br in list(el.iter(qn("w:br"))):
            if br.get(qn("w:type")) == "page":
                br.getparent().remove(br)
        for marker in list(el.iter(qn("w:lastRenderedPageBreak"))):
            marker.getparent().remove(marker)

    def _clean_runs(self, el, *, emphasis: bool, keep_italic: bool = True) -> None:
        for rpr in list(el.iter(qn("w:rPr"))):
            if rpr.getparent().tag == qn("w:pPr"):
                continue
            remove = list(_RUN_LOOK)
            if not emphasis:
                remove += [t for t in _RUN_EMPHASIS if keep_italic is False or t not in ("w:i", "w:iCs")]
            for tag in remove:
                for child in rpr.findall(qn(tag)):
                    rpr.remove(child)
            rstyle = rpr.find(qn("w:rStyle"))
            if rstyle is not None and rstyle.get(qn("w:val")) not in self.style_ids:
                rpr.remove(rstyle)

    # --- top level --------------------------------------------------------------------

    def render(self) -> bytes:
        blocks = [b for b in self.parsed.blocks if b.label != Label.IGNORE]
        title_page = [b for b in blocks if b.label in TITLE_PAGE_LABELS]
        abstract = [b for b in blocks if b.label in ABSTRACT_LABELS]
        body = [b for b in blocks if b.label not in TITLE_PAGE_LABELS and b.label not in ABSTRACT_LABELS]
        title = next((b for b in title_page if b.label == Label.TITLE), None)

        self._header(title)
        if self.options.include_title_page and title_page:
            self._title_page(title_page)
            self.pending_page_break = True
        if abstract:
            self._abstract(abstract)
            self.pending_page_break = True
        if title is not None:
            self._append(self._title_paragraph(title))
        self._body(body)

        buf = io.BytesIO()
        self.out.save(buf)
        return buf.getvalue()

    def _header(self, title: Block | None) -> None:
        section = self.out.sections[0]
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.style = self.out.styles["Header"]
        if self.options.paper_type == PaperType.PROFESSIONAL:
            running = self.options.running_head or (title.text if title else "")
            running = _shorten(running.upper(), 50)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(f"{running}\t")
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p._p.append(_page_field())

    def _title_paragraph(self, title: Block):
        el = self._copy_paragraph(title, "APA Title", emphasis=False)
        if self.options.title_case_headings:
            _rewrite_text(el, apa7.title_case)
        return el

    def _title_page(self, blocks: list[Block]) -> None:
        def of(label):
            return [b for b in blocks if b.label == label]

        for _ in range(3):
            self._append(self._new_paragraph("", "APA Title Page"))
        title = of(Label.TITLE)
        if title:
            self._append(self._title_paragraph(title[0]))
            self._append(self._new_paragraph("", "APA Title Page"))
        order = [Label.AUTHOR, Label.AFFILIATION]
        if self.options.paper_type == PaperType.STUDENT:
            order += [Label.COURSE, Label.INSTRUCTOR, Label.DATE]
        for label in order:
            for b in of(label):
                self._append(self._copy_paragraph(b, "APA Title Page", emphasis=False))
        notes = of(Label.AUTHOR_NOTE)
        if notes:
            self._append(self._new_paragraph("", "APA Title Page"))
            self._append(self._new_paragraph("Author Note", "APA Section Label"))
            for b in notes:
                self._append(self._copy_paragraph(b, "APA Body"))

    def _abstract(self, blocks: list[Block]) -> None:
        self._append(self._new_paragraph("Abstract", "APA Section Label"))
        for b in blocks:
            if b.label == Label.ABSTRACT_HEADING:
                continue
            if b.label == Label.KEYWORDS:
                words = apa7.strip_keywords_prefix(b.text).strip().rstrip(".")
                p = self._new_paragraph("Keywords:", "APA Keywords", italic=True)
                self._paragraph(p).add_run(f" {words}")
                self._append(p)
            else:
                el = self._copy_paragraph(b, "APA Abstract")
                _remove_prefix(el, len(paragraph_text(el)) - len(apa7.strip_abstract_prefix(paragraph_text(el))))
                self._append(el)

    def _body(self, blocks: list[Block]) -> None:
        blocks = _captions_first(blocks)
        references = [b for b in blocks if b.label == Label.REFERENCE_ENTRY]
        if self.options.sort_references:
            references.sort(key=lambda b: apa7.reference_sort_key(b.text))
        references_done = False
        last_label: Label | None = None

        i = 0
        while i < len(blocks):
            b = blocks[i]
            label = b.label
            nxt = blocks[i + 1] if i + 1 < len(blocks) else None

            if label in HEADING_LEVELS:
                level = HEADING_LEVELS[label]
                if level >= 4 and nxt is not None and nxt.label == Label.BODY:
                    self._append(self._run_in_heading(b, nxt, level))
                    i += 2
                    last_label = Label.BODY
                    continue
                self._append(self._heading(b, f"Heading {level}"))
            elif label in (Label.REFERENCE_HEADING, Label.REFERENCE_ENTRY):
                if not references_done:
                    self.pending_page_break = True
                    self._append(self._new_paragraph("References", "APA Section Label"))
                    for ref in references:
                        el = self._copy_paragraph(ref, "APA Reference")
                        text = paragraph_text(el)
                        _remove_prefix(el, len(text) - len(apa7.strip_reference_number(text)))
                        self._append(el)
                    references_done = True
            elif label == Label.APPENDIX_HEADING:
                if last_label != Label.APPENDIX_HEADING:
                    self.pending_page_break = True
                self._appendix_heading(b)
            elif label in (Label.FIGURE_CAPTION, Label.TABLE_CAPTION):
                self._caption(b)
            elif label == Label.FIGURE:
                self._append(self._copy_paragraph(b, "APA Figure"))
            elif label == Label.TABLE:
                self._append(self._table(b))
            elif label == Label.NOTE:
                self._note(b)
            elif label == Label.EQUATION:
                self._append(self._copy_paragraph(b, "APA Equation"))
            elif label == Label.BLOCK_QUOTE:
                self._append(self._copy_paragraph(b, "APA Block Quote"))
            elif label == Label.LIST_ITEM:
                self._append(self._copy_paragraph(b, "APA List", keep_numbering=True))
            elif label == Label.FOOTNOTE_TEXT:
                self._append(self._copy_paragraph(b, "APA Footnote"))
            else:
                self._append(self._copy_paragraph(b, "APA Body"))
            last_label = label
            i += 1

    # --- element builders -------------------------------------------------------------

    def _heading(self, b: Block, style: str):
        el = self._copy_paragraph(b, style, emphasis=False, keep_italic=not b.italic)
        text = paragraph_text(el)
        _remove_prefix(el, len(text) - len(apa7.strip_heading_number(text)))
        if self.options.title_case_headings:
            _rewrite_text(el, apa7.title_case)
        return el

    def _run_in_heading(self, heading: Block, body: Block, level: int):
        head = self._heading(heading, "APA Body")
        text = paragraph_text(head).rstrip()
        para = self._copy_paragraph(body, "APA Body")
        ppr = para.find(qn("w:pPr"))
        insert_at = 1 if ppr is not None else 0
        runs = [c for c in head if c.tag != qn("w:pPr")]
        if not text.endswith((".", "?", "!", ":")):
            runs.append(_make_run("."))
        runs.append(_make_run(" ", bold=False))
        for r in runs:
            for run_el in ([r] if r.tag == qn("w:r") else r.iter(qn("w:r"))):
                if run_el.find(qn("w:t")) is not None and paragraph_text(run_el).strip():
                    Run(run_el, None).bold = True
                    if level == 5:
                        Run(run_el, None).italic = True
            para.insert(insert_at, r)
            insert_at += 1
        return para

    def _appendix_heading(self, b: Block) -> None:
        text = b.text
        m = re.match(r"^(appendix(\s+[A-Z0-9]{1,3})?)\s*[:.\-–—]\s*(\S.*)$", text, re.I)
        if m:
            self._append(self._new_paragraph(m.group(1).strip().title(), "APA Section Label"))
            title = m.group(3).strip()
            if self.options.title_case_headings:
                title = apa7.title_case(title)
            self._append(self._new_paragraph(title, "APA Section Label"))
        else:
            el = self._copy_paragraph(b, "APA Section Label", emphasis=False)
            if self.options.title_case_headings:
                _rewrite_text(el, apa7.title_case)
            self._append(el)

    def _caption(self, b: Block) -> None:
        parts = apa7.split_caption(b.text)
        if parts is None:
            el = self._copy_paragraph(b, "APA Caption Title", emphasis=False)
            if self.options.title_case_headings:
                _rewrite_text(el, apa7.title_case)
            self._append(el)
            return
        number, title = parts
        self._append(self._new_paragraph(number, "APA Caption Label"))
        if title:
            el = self._copy_paragraph(b, "APA Caption Title", emphasis=False, keep_italic=False)
            _remove_prefix(el, apa7.caption_prefix_length(paragraph_text(el)))
            _rewrite_text(el, lambda t: (apa7.title_case(t) if self.options.title_case_headings else t).rstrip("."))
            self._append(el)

    def _note(self, b: Block) -> None:
        el = self._copy_paragraph(b, "APA Note")
        text = paragraph_text(el)
        m = re.match(r"^\s*notes?\s*[.:]\s*", text, re.I)
        if m:
            _remove_prefix(el, m.end())
        idx = 1 if el.find(qn("w:pPr")) is not None else 0
        el.insert(idx, _make_run("Note.", italic=True))
        el.insert(idx + 1, _make_run(" "))
        self._append(el)

    def _table(self, b: Block):
        tbl = self.copier.copy(self.parsed.elements[b.index])
        tbl_pr = tbl.find(qn("w:tblPr"))
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            tbl.insert(0, tbl_pr)
        for tag in ("w:tblStyle", "w:tblBorders", "w:shd", "w:tblLook"):
            for child in tbl_pr.findall(qn(tag)):
                tbl_pr.remove(child)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement(f"w:{side}")
            if side in ("top", "bottom"):
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), "8")
                edge.set(qn("w:space"), "0")
                edge.set(qn("w:color"), "000000")
            else:
                edge.set(qn("w:val"), "nil")
            borders.append(edge)
        _insert_in_order(tbl_pr, borders, ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
                                           "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
                                           "w:tblCellSpacing", "w:tblInd"))
        for tc_pr in tbl.iter(qn("w:tcPr")):
            for tag in ("w:tcBorders", "w:shd"):
                for child in tc_pr.findall(qn(tag)):
                    tc_pr.remove(child)
        rows = tbl.findall(qn("w:tr"))
        if rows:
            header = rows[0]
            tr_pr = header.find(qn("w:trPr"))
            if tr_pr is None:
                tr_pr = OxmlElement("w:trPr")
                tbl_pr_ex = header.find(qn("w:tblPrEx"))
                header.insert(1 if tbl_pr_ex is not None else 0, tr_pr)
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
            for tc in header.findall(qn("w:tc")):
                tc_pr = tc.find(qn("w:tcPr"))
                if tc_pr is None:
                    tc_pr = OxmlElement("w:tcPr")
                    tc.insert(0, tc_pr)
                tc_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "0"), ("w:color", "000000")):
                    bottom.set(qn(k), v)
                tc_borders.append(bottom)
                _insert_in_order(tc_pr, tc_borders, ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge"))
        for p in tbl.iter(qn("w:p")):
            self._clean_paragraph(p, "APA Table Text", keep_alignment=True)
            self._clean_runs(p, emphasis=True)
        return tbl


# --- XML helpers --------------------------------------------------------------------------


def _inline_style_numbering(src, el, src_doc) -> None:
    """Lists numbered through their style ("List Number") lose that style; copy its numPr inline."""
    if el.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None:
        return
    try:
        style = Paragraph(src, src_doc).style
    except (KeyError, ValueError):
        return
    while style is not None:
        num_pr = style.element.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if num_pr is not None and num_pr.find(qn("w:numId")) is not None:
            new = copy.deepcopy(num_pr)
            if new.find(qn("w:ilvl")) is None:
                ilvl = OxmlElement("w:ilvl")
                ilvl.set(qn("w:val"), "0")
                new.insert(0, ilvl)
            Paragraph(el, None)._p.get_or_add_pPr()._insert_numPr(new)
            return
        style = style.base_style


def _text_nodes(el):
    for node in el.iter(qn("w:t"), qn("w:tab"), qn("w:br")):
        if node.tag == qn("w:br") and node.get(qn("w:type")) in ("page", "column"):
            continue
        if node.tag == qn("w:tab") and node.getparent().tag != qn("w:r"):
            continue  # tab stop definitions, not tab characters
        yield node


def _remove_prefix(el, n: int) -> None:
    """Delete the first n characters of a paragraph's text, across runs."""
    for node in list(_text_nodes(el)):
        if n <= 0:
            return
        if node.tag == qn("w:t"):
            text = node.text or ""
            take = min(n, len(text))
            node.text = text[take:]
            n -= take
        else:
            node.getparent().remove(node)
            n -= 1


def _strip_leading_whitespace(el) -> None:
    text = paragraph_text(el)
    _remove_prefix(el, len(text) - len(text.lstrip()))


def _rewrite_text(el, fn) -> None:
    """Apply a length-preserving text transform (e.g. title case) across runs."""
    nodes = [n for n in el.iter(qn("w:t"))]
    original = "".join(n.text or "" for n in nodes)
    new = fn(original)
    if len(new) != len(original):
        # Only safe to redistribute when lengths match; otherwise trim trailing characters.
        if new == original[: len(new)]:
            _remove_suffix(nodes, len(original) - len(new))
        return
    pos = 0
    for n in nodes:
        length = len(n.text or "")
        n.text = new[pos:pos + length]
        pos += length


def _remove_suffix(nodes, n: int) -> None:
    for node in reversed(nodes):
        if n <= 0:
            return
        text = node.text or ""
        take = min(n, len(text))
        node.text = text[: len(text) - take]
        n -= take


def _make_run(text: str, *, bold: bool | None = None, italic: bool | None = None):
    r = OxmlElement("w:r")
    run = Run(r, None)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def _page_field():
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fld.append(_make_run("1"))
    return fld


def _insert_in_order(parent, child, predecessors: tuple[str, ...]) -> None:
    """Insert child after the last existing element whose tag is in predecessors."""
    tags = {qn(t) for t in predecessors}
    idx = 0
    for i, existing in enumerate(parent):
        if existing.tag in tags:
            idx = i + 1
    parent.insert(idx, child)


def _shorten(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0]
    return cut.rstrip(" :;,-")


def _captions_first(blocks: list[Block]) -> list[Block]:
    """APA puts the 'Figure N' / 'Table N' caption above its figure or table."""
    out = list(blocks)
    i = 0
    while i < len(out) - 1:
        cur, nxt = out[i], out[i + 1]
        if (cur.label == Label.FIGURE and nxt.label == Label.FIGURE_CAPTION) or (
            cur.label == Label.TABLE and nxt.label == Label.TABLE_CAPTION
        ):
            out[i], out[i + 1] = nxt, cur
            i += 2
        else:
            i += 1
    return out


def render_docx(parsed: ParsedDocument, options: RenderOptions) -> bytes:
    return _Renderer(parsed, options).render()
